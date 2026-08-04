from sqlalchemy.orm import Session
from sqlalchemy import or_
from app.models import IndicatorLibrary
from app.schemas.indicator_library import (
    IndicatorLibraryCreate,
    BatchCreateItem,
    BatchUpdateItem,
)
from app.crud.indicator_library_version import create_version_from_indicator, delete_versions_for_indicator
import json


def get_library_indicators(
    db: Session,
    skip: int = 0,
    limit: int = 100,
    search: str = None,
    category: str = None,
    data_type: str = None,
    is_required: bool = None,
):
    """Получить показатели из справочника с фильтрацией и поиском."""
    query = db.query(IndicatorLibrary)

    # Поиск по названию, описанию и категории
    if search:
        search_term = f"%{search}%"
        query = query.filter(
            or_(
                IndicatorLibrary.name.ilike(search_term),
                IndicatorLibrary.description.ilike(search_term),
                IndicatorLibrary.category.ilike(search_term),
            )
        )

    # Фильтрация по категории
    if category:
        query = query.filter(IndicatorLibrary.category == category)

    # Фильтрация по типу данных
    if data_type:
        query = query.filter(IndicatorLibrary.data_type == data_type)

    # Фильтрация по обязательности
    if is_required is not None:
        query = query.filter(IndicatorLibrary.is_required == is_required)

    # Сортировка по id (самые новые в конце)
    query = query.order_by(IndicatorLibrary.id.asc())

    return query.offset(skip).limit(limit).all()


def get_library_indicators_count(
    db: Session,
    search: str = None,
    category: str = None,
    data_type: str = None,
    is_required: bool = None,
) -> int:
    """Получить общее количество показателей по фильтру (для пагинации)."""
    query = db.query(IndicatorLibrary)

    if search:
        search_term = f"%{search}%"
        query = query.filter(
            or_(
                IndicatorLibrary.name.ilike(search_term),
                IndicatorLibrary.description.ilike(search_term),
                IndicatorLibrary.category.ilike(search_term),
            )
        )

    if category:
        query = query.filter(IndicatorLibrary.category == category)

    if data_type:
        query = query.filter(IndicatorLibrary.data_type == data_type)

    if is_required is not None:
        query = query.filter(IndicatorLibrary.is_required == is_required)

    return query.count()


def get_library_indicator(db: Session, indicator_id: int):
    return db.query(IndicatorLibrary).filter(IndicatorLibrary.id == indicator_id).first()


def get_library_indicator_by_name(db: Session, name: str):
    return db.query(IndicatorLibrary).filter(IndicatorLibrary.name == name).first()


def create_library_indicator(db: Session, indicator: IndicatorLibraryCreate, user_id: int = None):
    # Конвертируем options в JSON-строку
    options_json = None
    if indicator.options and indicator.data_type == 'select':
        options_json = json.dumps(indicator.options)

    # Конвертируем validation_rules в JSON-строку, если передан
    validation_json = None
    if indicator.validation_rules:
        import json as json_module
        if isinstance(indicator.validation_rules, (dict, list)):
            validation_json = json_module.dumps(indicator.validation_rules)
        else:
            validation_json = indicator.validation_rules

    db_indicator = IndicatorLibrary(
        name=indicator.name,
        unit=indicator.unit,
        data_type=indicator.data_type,
        options=options_json,
        description=indicator.description,
        category=indicator.category,
        is_required=indicator.is_required,
        default_value=indicator.default_value,
        validation_rules=validation_json,
        created_by=user_id,
    )
    db.add(db_indicator)
    db.commit()
    db.refresh(db_indicator)

    # Создаём первую версию (create)
    create_version_from_indicator(
        db, db_indicator,
        change_type="create",
        changed_by=user_id,
        change_notes="Показатель создан"
    )

    return db_indicator


def update_library_indicator(db: Session, indicator_id: int, indicator: IndicatorLibraryCreate, user_id: int = None):
    db_indicator = get_library_indicator(db, indicator_id)
    if not db_indicator:
        return None

    options_json = None
    if indicator.options and indicator.data_type == 'select':
        options_json = json.dumps(indicator.options)

    # Сохраняем старый name для change_notes
    old_name = db_indicator.name

    db_indicator.name = indicator.name
    db_indicator.unit = indicator.unit
    db_indicator.data_type = indicator.data_type
    db_indicator.options = options_json
    db_indicator.description = indicator.description
    db_indicator.category = indicator.category
    db_indicator.is_required = indicator.is_required
    db_indicator.default_value = indicator.default_value

    # Конвертируем validation_rules в JSON-строку, если передан
    if indicator.validation_rules:
        import json as json_module
        if isinstance(indicator.validation_rules, (dict, list)):
            db_indicator.validation_rules = json_module.dumps(indicator.validation_rules)
        else:
            db_indicator.validation_rules = indicator.validation_rules
    else:
        db_indicator.validation_rules = None

    db.commit()
    db.refresh(db_indicator)

    # Создаём версию изменения
    create_version_from_indicator(
        db, db_indicator,
        change_type="update",
        changed_by=user_id,
        change_notes=f"Обновление показателя: {old_name}"
    )

    return db_indicator


def delete_library_indicator(db: Session, indicator_id: int, user_id: int = None):
    db_indicator = get_library_indicator(db, indicator_id)
    if db_indicator:
        # Удаляем все версии, связанные с показателем
        delete_versions_for_indicator(db, indicator_id)
        
        # Удаляем сам показатель (каскадное удаление версий уже не нужно)
        db.delete(db_indicator)
        db.commit()
    return db_indicator


# ---- Batch операции ----

def batch_create_indicators(
    db: Session,
    indicators: list[BatchCreateItem],
    user_id: int = None,
) -> tuple[list[IndicatorLibrary], list[dict]]:
    """Создать несколько показателей за один запрос."""
    created = []
    errors = []

    for item in indicators:
        try:
            # Проверка на дубликат
            existing = get_library_indicator_by_name(db, item.name)
            if existing:
                errors.append({
                    "name": item.name,
                    "error": f"Показатель '{item.name}' уже существует",
                })
                continue

            # Создаём через базовую схему
            create_data = IndicatorLibraryCreate(
                name=item.name,
                unit=item.unit,
                data_type=item.data_type,
                options=item.options,
                description=item.description,
                category=item.category,
                is_required=item.is_required,
                default_value=item.default_value,
                validation_rules=item.validation_rules,
            )
            db_indicator = create_library_indicator(db, create_data, user_id=user_id)
            created.append(db_indicator)
        except Exception as e:
            errors.append({
                "name": item.name,
                "error": str(e),
            })

    return created, errors


def batch_update_indicators(
    db: Session,
    indicators: list[BatchUpdateItem],
    user_id: int = None,
) -> tuple[list[IndicatorLibrary], list[dict]]:
    """Обновить несколько показателей за один запрос."""
    updated = []
    errors = []

    for item in indicators:
        try:
            db_indicator = get_library_indicator(db, item.id)
            if not db_indicator:
                errors.append({
                    "id": item.id,
                    "error": f"Показатель с id={item.id} не найден",
                })
                continue

            # Обновляем только переданные поля
            old_name = db_indicator.name

            if item.name is not None:
                # Проверка на дубликат имени (если имя меняется)
                if item.name != db_indicator.name:
                    existing = get_library_indicator_by_name(db, item.name)
                    if existing:
                        errors.append({
                            "id": item.id,
                            "error": f"Показатель с именем '{item.name}' уже существует",
                        })
                        continue
                db_indicator.name = item.name
            if item.unit is not None:
                db_indicator.unit = item.unit
            if item.data_type is not None:
                db_indicator.data_type = item.data_type
            if item.options is not None:
                if item.data_type == 'select' or db_indicator.data_type == 'select':
                    db_indicator.options = json.dumps(item.options)
                else:
                    db_indicator.options = None
            if item.description is not None:
                db_indicator.description = item.description
            if item.category is not None:
                db_indicator.category = item.category
            if item.is_required is not None:
                db_indicator.is_required = item.is_required
            if item.default_value is not None:
                db_indicator.default_value = item.default_value
            if item.validation_rules is not None:
                db_indicator.validation_rules = item.validation_rules

            db.commit()
            db.refresh(db_indicator)

            # Создаём версию изменения
            create_version_from_indicator(
                db, db_indicator,
                change_type="update",
                changed_by=user_id,
                change_notes=f"Batch-обновление показателя: {old_name}"
            )

            updated.append(db_indicator)
        except Exception as e:
            errors.append({
                "id": item.id,
                "error": str(e),
            })

    return updated, errors


def export_indicators_csv(db: Session) -> str:
    """Экспорт справочника показателей в CSV (строку)."""
    import csv
    import io

    indicators = db.query(IndicatorLibrary).order_by(IndicatorLibrary.id).all()
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Заголовки
    writer.writerow([
        'id', 'name', 'unit', 'data_type', 'options', 
        'description', 'category', 'is_required', 'default_value', 'validation_rules'
    ])
    
    for ind in indicators:
        options_str = ''
        if ind.options:
            try:
                opts = json.loads(ind.options)
                options_str = ';'.join(opts)
            except (json.JSONDecodeError, TypeError):
                options_str = ind.options or ''
        
        writer.writerow([
            ind.id,
            ind.name,
            ind.unit or '',
            ind.data_type or 'number',
            options_str,
            ind.description or '',
            ind.category or '',
            '1' if ind.is_required else '0',
            ind.default_value or '',
            ind.validation_rules or '',
        ])
    
    return output.getvalue()


def export_indicators_excel(db: Session) -> bytes:
    """Экспорт справочника показателей в Excel (.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    indicators = db.query(IndicatorLibrary).order_by(IndicatorLibrary.id).all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Справочник показателей"
    
    # Заголовки
    headers = [
        'ID', 'Название', 'Ед. изм.', 'Тип данных', 'Варианты (;)',
        'Описание', 'Категория', 'Обязательный', 'Значение по умолчанию', 'Правила валидации'
    ]
    
    # Стили
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    
    # Данные
    type_labels = {'number': 'Число', 'text': 'Текст', 'select': 'Выбор'}
    category_labels = {
        'quality': 'Качество', 'safety': 'Безопасность', 'performance': 'Производительность',
        'chemical': 'Химический состав', 'physical': 'Физические свойства', 'microbiology': 'Микробиология',
    }
    
    for row_idx, ind in enumerate(indicators, 2):
        options_str = ''
        if ind.options:
            try:
                opts = json.loads(ind.options)
                options_str = '; '.join(opts)
            except (json.JSONDecodeError, TypeError):
                options_str = ind.options or ''
        
        ws.cell(row=row_idx, column=1, value=ind.id).border = thin_border
        ws.cell(row=row_idx, column=2, value=ind.name).border = thin_border
        ws.cell(row=row_idx, column=3, value=ind.unit or '').border = thin_border
        ws.cell(row=row_idx, column=4, value=type_labels.get(ind.data_type, ind.data_type)).border = thin_border
        ws.cell(row=row_idx, column=5, value=options_str).border = thin_border
        ws.cell(row=row_idx, column=6, value=ind.description or '').border = thin_border
        ws.cell(row=row_idx, column=7, value=category_labels.get(ind.category, ind.category or '')).border = thin_border
        ws.cell(row=row_idx, column=8, value='Да' if ind.is_required else 'Нет').border = thin_border
        ws.cell(row=row_idx, column=9, value=ind.default_value or '').border = thin_border
        ws.cell(row=row_idx, column=10, value=ind.validation_rules or '').border = thin_border
    
    # Автоширина колонок
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 60)
    
    import io
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def import_indicators_csv(
    db: Session,
    csv_content: str,
    user_id: int = None,
) -> tuple[list[dict], list[dict]]:
    """Импорт показателей из CSV. Возвращает (created, errors)."""
    import csv
    import io

    created = []
    errors = []

    reader = csv.DictReader(io.StringIO(csv_content))
    
    # Маппинг русских заголовков на английские (для Excel-экспорта)
    header_map = {
        'Название': 'name',
        'Ед. изм.': 'unit',
        'Тип данных': 'data_type',
        'Варианты (;)': 'options',
        'Описание': 'description',
        'Категория': 'category',
        'Обязательный': 'is_required',
        'Значение по умолчанию': 'default_value',
        'Правила валидации': 'validation_rules',
    }
    
    for row_idx, row in enumerate(reader, 2):
        try:
            # Нормализуем заголовки
            name = row.get('name') or row.get('Название', '').strip()
            if not name:
                errors.append({"row": row_idx, "error": "Пустое название"})
                continue
            
            # Проверка дубликатов
            existing = get_library_indicator_by_name(db, name)
            if existing:
                errors.append({"row": row_idx, "error": f"Показатель '{name}' уже существует", "name": name})
                continue
            
            unit = row.get('unit') or row.get('Ед. изм.', '').strip()
            data_type_raw = (row.get('data_type') or row.get('Тип данных', 'number')).strip().lower()
            
            # Маппинг русских названий типов
            type_map = {'число': 'number', 'текст': 'text', 'выбор': 'select'}
            data_type = type_map.get(data_type_raw, data_type_raw)
            if data_type not in ('number', 'text', 'select'):
                data_type = 'number'
            
            options_raw = row.get('options') or row.get('Варианты (;)', '').strip()
            options = None
            if options_raw and data_type == 'select':
                options = [x.strip() for x in options_raw.replace('; ', ';').split(';') if x.strip()]
            
            description = row.get('description') or row.get('Описание', '').strip() or None
            category_raw = (row.get('category') or row.get('Категория', '')).strip()
            
            # Маппинг русских категорий
            category_map = {
                'качество': 'quality', 'безопасность': 'safety', 'производительность': 'performance',
                'химический состав': 'chemical', 'физические свойства': 'physical', 'микробиология': 'microbiology',
            }
            category = category_map.get(category_raw.lower(), category_raw) if category_raw else None
            
            is_required_raw = (row.get('is_required') or row.get('Обязательный', '0')).strip().lower()
            is_required = is_required_raw in ('1', 'yes', 'да', 'true')
            
            default_value = row.get('default_value') or row.get('Значение по умолчанию', '').strip() or None
            validation_rules = row.get('validation_rules') or row.get('Правила валидации', '').strip() or None
            
            create_data = IndicatorLibraryCreate(
                name=name,
                unit=unit,
                data_type=data_type,
                options=options,
                description=description,
                category=category,
                is_required=is_required,
                default_value=default_value,
                validation_rules=validation_rules,
            )
            db_indicator = create_library_indicator(db, create_data, user_id=user_id)
            created.append({
                "id": db_indicator.id,
                "name": db_indicator.name,
            })
        except Exception as e:
            errors.append({"row": row_idx, "error": str(e), "name": row.get('name', '')})
    
    return created, errors


def batch_delete_indicators(
    db: Session,
    ids: list[int],
    user_id: int = None,
) -> tuple[list[int], list[dict]]:
    """Удалить несколько показателей за один запрос."""
    deleted_ids = []
    errors = []

    for indicator_id in ids:
        try:
            db_indicator = delete_library_indicator(db, indicator_id, user_id=user_id)
            if db_indicator:
                deleted_ids.append(indicator_id)
            else:
                errors.append({
                    "id": indicator_id,
                    "error": f"Показатель с id={indicator_id} не найден",
                })
        except Exception as e:
            errors.append({
                "id": indicator_id,
                "error": str(e),
            })

    return deleted_ids, errors