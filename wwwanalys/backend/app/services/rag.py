"""
База знаний технологических карт (RAG).

Пайплайн: файл → извлечение текста → нарезка на фрагменты → хранение в Postgres
(таблицы tech_documents/tech_chunks) в вектор-готовом виде.

Поиск сейчас идёт по полнотекстовому индексу Postgres (русская конфигурация) —
работает без эмбеддинг-модели. Когда модель выбрана, фрагменты до-индексируются
векторами (pgvector), и включается семантический/гибридный поиск — аддитивно.

Эмбеддинг-слой спрятан за EmbeddingProvider (по аналогии с экспертным слоем),
поэтому конкретную модель можно подключить позже, не меняя пайплайн и API.
"""
import hashlib
import io
import logging
import re
from typing import Any, Dict, List, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import Base, engine
from app.models.tech_card import TechDocument, TechChunk

logger = logging.getLogger(__name__)


# ==================== Инициализация хранилища ====================

def init_storage() -> None:
    """Создать расширение pgvector, таблицы и индексы. Идемпотентно.

    Вызывается при загрузке AI-модуля в защищённом блоке — исключение здесь
    не должно ронять основное приложение (обрабатывается в вызывающем коде).
    """
    with engine.connect() as conn:
        conn.execute(text("CREATE EXTENSION IF NOT EXISTS vector"))
        conn.commit()

    Base.metadata.create_all(
        bind=engine, tables=[TechDocument.__table__, TechChunk.__table__]
    )

    with engine.connect() as conn:
        # Полнотекстовый поиск (русский) — работает уже сейчас, без эмбеддингов.
        conn.execute(text(
            "CREATE INDEX IF NOT EXISTS ix_tech_chunks_fts "
            "ON tech_chunks USING gin (to_tsvector('russian', content))"
        ))
        # Векторный индекс (косинус). На пустой/NULL-колонке безвреден.
        try:
            conn.execute(text(
                "CREATE INDEX IF NOT EXISTS ix_tech_chunks_embedding "
                "ON tech_chunks USING hnsw (embedding vector_cosine_ops)"
            ))
        except Exception as e:  # некоторые сборки pgvector без hnsw — не критично
            logger.warning("HNSW-индекс не создан (%s); поиск по вектору без ANN", e)
        conn.commit()
    logger.info("RAG storage initialised")


# ==================== Извлечение текста из файлов ====================

def extract_text(filename: str, data: bytes, mime: Optional[str] = None) -> str:
    """Достать текст из файла по расширению. Поддержка: pdf, docx, xlsx, txt, md, csv."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".xlsx"):
        return _extract_xlsx(data)
    if name.endswith((".txt", ".md", ".csv")):
        return _decode_text(data)
    # запасной вариант — попытаться как текст
    return _decode_text(data)


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # таблицы
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# ==================== Нарезка на фрагменты ====================

def chunk_text(
    text_content: str,
    chunk_size: Optional[int] = None,
    overlap: Optional[int] = None,
) -> List[str]:
    """Разбить текст на фрагменты по абзацам с учётом размера и перекрытия."""
    size = chunk_size or settings.rag_chunk_size
    ov = overlap if overlap is not None else settings.rag_chunk_overlap

    # нормализуем пробелы/переводы строк, режем на абзацы
    paragraphs = re.split(r"\n\s*\n", text_content.strip())
    paragraphs = [re.sub(r"[ \t]+", " ", p).strip() for p in paragraphs if p.strip()]

    chunks: List[str] = []
    buf = ""
    for p in paragraphs:
        if len(p) > size:
            # длинный абзац — режем по предложениям/жёстко
            if buf:
                chunks.append(buf)
                buf = ""
            for piece in _split_long(p, size):
                chunks.append(piece)
            continue
        if len(buf) + len(p) + 1 <= size:
            buf = f"{buf}\n{p}" if buf else p
        else:
            chunks.append(buf)
            buf = p
    if buf:
        chunks.append(buf)

    # добавляем перекрытие (хвост предыдущего фрагмента в начало следующего)
    if ov > 0 and len(chunks) > 1:
        with_ov: List[str] = []
        for i, c in enumerate(chunks):
            if i == 0:
                with_ov.append(c)
            else:
                tail = chunks[i - 1][-ov:]
                with_ov.append(f"{tail}\n{c}")
        chunks = with_ov
    return chunks


def _split_long(paragraph: str, size: int) -> List[str]:
    sentences = re.split(r"(?<=[.!?])\s+", paragraph)
    out: List[str] = []
    buf = ""
    for s in sentences:
        if len(s) > size:
            if buf:
                out.append(buf)
                buf = ""
            for i in range(0, len(s), size):
                out.append(s[i:i + size])
            continue
        if len(buf) + len(s) + 1 <= size:
            buf = f"{buf} {s}" if buf else s
        else:
            out.append(buf)
            buf = s
    if buf:
        out.append(buf)
    return out


# ==================== Загрузка документа ====================

def ingest_document(
    db: Session,
    *,
    title: str,
    filename: str,
    data: bytes,
    mime: Optional[str],
    variety: Optional[str],
    uploaded_by: Optional[int],
) -> TechDocument:
    """Разобрать файл, нарезать и сохранить документ с фрагментами."""
    sha = hashlib.sha256(data).hexdigest()

    raw = extract_text(filename, data, mime)
    pieces = chunk_text(raw)

    doc = TechDocument(
        title=title or filename,
        variety=variety,
        filename=filename,
        mime=mime,
        size_bytes=len(data),
        sha256=sha,
        status="ready",
        char_count=len(raw),
        chunk_count=len(pieces),
        uploaded_by=uploaded_by,
    )
    db.add(doc)
    db.flush()  # получить doc.id

    for i, content in enumerate(pieces):
        db.add(TechChunk(
            document_id=doc.id,
            chunk_index=i,
            content=content,
            char_count=len(content),
            meta={"variety": variety} if variety else None,
            embedding=None,  # заполнится при подключении эмбеддинг-модели
        ))
    db.commit()
    db.refresh(doc)
    return doc


# ==================== Поиск (FTS сейчас, вектор — позже) ====================

def search_chunks(
    db: Session,
    query: str,
    *,
    variety: Optional[str] = None,
    limit: int = 5,
) -> List[Dict[str, Any]]:
    """Поиск релевантных фрагментов. Пока — полнотекстовый (русский)."""
    params: Dict[str, Any] = {"q": query, "limit": limit}
    variety_clause = ""
    if variety:
        variety_clause = "AND (d.variety = :variety OR d.variety IS NULL)"
        params["variety"] = variety

    sql = text(f"""
        SELECT c.id, c.document_id, c.content, d.title, d.variety,
               ts_rank(to_tsvector('russian', c.content),
                       plainto_tsquery('russian', :q)) AS rank
        FROM tech_chunks c
        JOIN tech_documents d ON d.id = c.document_id
        WHERE to_tsvector('russian', c.content) @@ plainto_tsquery('russian', :q)
        {variety_clause}
        ORDER BY rank DESC
        LIMIT :limit
    """)
    rows = db.execute(sql, params).fetchall()
    return [
        {
            "chunk_id": r.id,
            "document_id": r.document_id,
            "title": r.title,
            "variety": r.variety,
            "content": r.content,
            "score": float(r.rank),
        }
        for r in rows
    ]


def is_embedding_configured() -> bool:
    return (settings.rag_embedding_provider or "none").strip().lower() != "none"
